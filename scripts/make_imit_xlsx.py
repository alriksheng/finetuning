"""仿写样本抽查表：每条意见一行，供人工判定（可用/需修改/不可用）。
Usage: python3 -I make_imit_xlsx.py <data_dir> <批次号> <输出xlsx> <标注文件.py ...>  （只列这些标注文件里的样本）"""
import sys, os, re, json, runpy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
D, NO, OUT = sys.argv[1], sys.argv[2], sys.argv[3]; FILES = sys.argv[4:]
want = []
for f in FILES:
    for it in runpy.run_path(f)['ITEMS']: want.append('V7_' + it['src'][2:] + it.get('suffix', ''))
S = {json.loads(l)['id']: json.loads(l) for l in open(os.path.join(D, os.environ.get('IMIT_DIR', 'out_imit'), 'samples.jsonl'), encoding='utf-8')}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '抽查'
title = f'仿写样本抽查（第 {NO} 批，{len(want)} 个样本）：输入是真实报告小节，意见由我按评审人口吻撰写；“人造错”是对真实小节做了最小改动。请判断每条意见：可用=具体、成立、写法像评审人；需修改=基本成立但写法/表述要改；不可用=不成立或太泛。批注请写在备注列。'
ws.merge_cells('A1:J1'); ws['A1'] = title; ws['A1'].alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[1].height = 48
H = ['序号', '样本', '项目', '来源', '规则', '意见（问题+建议）', '引文', '思考（节选）', '判定', '备注']
fnt = 'Arial'
for c, h in enumerate(H, 1):
    x = ws.cell(2, c, h); x.font = Font(name=fnt, bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78'); x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin = Side(style='thin', color='BBBBBB'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
r = 3; n = 0
for sid in want:
    rec = S[sid]; asst = rec['messages'][1]['content']; th, js = asst.split('</think>', 1)
    th = th.replace('<think>', '').strip(); F = json.loads(js)['findings']
    src = '人造错（改过原文）' if rec['meta'].get('edits') else '真实小节'
    for f in F:
        n += 1
        vals = [n, sid, rec['meta']['project'], src, f['rule_id'], f['problem'] + f['suggestion'], '\n'.join(e['quote'] for e in f['evidence']), th[:420] + ('……' if len(th) > 420 else ''), None, None]
        for c, v in enumerate(vals, 1):
            x = ws.cell(r, c, v); x.font = Font(name=fnt, size=10); x.alignment = Alignment(wrap_text=True, vertical='top'); x.border = bd
        r += 1
for k, w in zip('ABCDEFGHIJ', [6, 18, 22, 14, 16, 56, 48, 60, 12, 24]): ws.column_dimensions[k].width = w
dv = DataValidation(type='list', formula1='"可用,需修改,不可用"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'I3:I{r - 1}')
ws.freeze_panes = 'C3'
wb.save(OUT); print(OUT, n, '行')
