"""开发集额外意见判定表：e3 单次采样里没有对上任何专家意见的意见，供专家判定“成立与否”。
Usage: python3 -I make_extras_xlsx.py <judge.jsonl> <inputs.jsonl> <out.xlsx>"""
import sys, json, re, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
jd, inp, out = sys.argv[1:4]
I = {r['id']: r for r in map(json.loads, open(inp, encoding='utf-8'))}
J = [json.loads(l) for l in open(jd, encoding='utf-8')]
parse = lambda t: json.loads(t.split('</think>')[-1].strip())
ftxt = lambda f: ((f.get('problem') or '') + (f.get('suggestion') or '')).strip()
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '待判定'
ws['A1'] = ('开发集额外意见判定：以下是模型（e3 单次采样）提出、但没有对上任何专家意见的意见。请对照右侧“同一小节专家已提的意见”和模型引文，判断这条意见本身是否成立。'
            '成立=报告确有此问题，评审中应当提出；不成立=报告没有此问题或模型理解有误；成立但次要=属实但一般不会写进评审意见；说不清=需要更多材料。')
ws.merge_cells('A1:K1'); ws['A1'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[1].height = 60
hdr = ['序号', '样本', '项目', '小节', '规则', '模型意见', '模型引文', '同一小节专家已提的意见', '判定', '备注', '模型引文位置']
ws.append(hdr)
thin = Side(style='thin', color='BBBBBB')
for c in ws[2]:
    c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
rows = []
for n, x in enumerate(J, 1):
    r = I[x['id']]; f = x['finding']; gold = parse(r['gold'])['findings']
    sec = re.search(r'\[([^\]]+)\]', r['prompt'].split('【报告段落】')[-1]); sec = sec.group(1) if sec else ''
    gtxt = '\n'.join(f"· [{g['rule_id']}] {ftxt(g)}" for g in gold)
    q = '\n'.join(e.get('quote', '') for e in f.get('evidence') or []); loc = '；'.join(e.get('loc', '') for e in f.get('evidence') or [])
    ws.append([n, x['id'], x['project'], sec[:40], f.get('rule_id'), ftxt(f), q, gtxt, None, None, loc])
widths = [6, 16, 8, 24, 18, 50, 50, 60, 14, 24, 24]
for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    row[8].fill = PatternFill('solid', fgColor='FFF2CC')
dv = DataValidation(type='list', formula1='"成立,不成立,成立但次要,说不清"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'I3:I{ws.max_row}')
ws.freeze_panes = 'C3'; wb.save(out); print(len(J), '条 →', out)
