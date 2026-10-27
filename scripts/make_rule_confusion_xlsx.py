"""开发集“位置找对、规则贴错”清单：对每条标准意见，若某条预测意见的引文覆盖了标准引文（口径同 eval_score.py），但没有任何命中引文的预测意见与标准规则相同，则列出。
Usage: python3 -I make_rule_confusion_xlsx.py <inputs.jsonl> <pred.jsonl> <rules.yaml> <out.xlsx>"""
import sys, json, re, difflib, collections, yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
INP, PRED, RULES, OUT = sys.argv[1:5]
R = {r['rule_id']: r for r in yaml.safe_load(open(RULES, encoding='utf-8'))}
squash = lambda s: re.sub(r'\s+', '', s or '')
def qh(g, p):
    g, p = squash(g), squash(p)
    if not g or not p: return False
    m = difflib.SequenceMatcher(None, g, p, autojunk=False).find_longest_match(0, len(g), 0, len(p))
    return m.size >= 10 and m.size >= 0.5 * min(len(g), len(p))
def parse(t):
    t = (t or '').split('</think>')[-1]; i = t.find('{')
    try: return json.JSONDecoder().raw_decode(t[i:])[0]
    except Exception: return None
Q = lambda f: [e.get('quote', '') for e in (f.get('evidence') or []) if isinstance(e, dict)]
txt = lambda f: (f.get('problem') or '') + (f.get('suggestion') or '')
name = lambda r: f"{r}｜{R[r]['tag']}" if r in R else str(r)
pts = lambda r: (R[r].get('question', '') + '。' + R[r].get('judgment_points', '')) if r in R else ''
inp = {json.loads(l)['id']: json.loads(l) for l in open(INP, encoding='utf-8')}
pred = {json.loads(l)['id']: json.loads(l) for l in open(PRED, encoding='utf-8')}
rows = []
for i, r in inp.items():
    out = parse(pred[i]['output']) if i in pred else None
    F = [f for f in (out or {}).get('findings', []) if isinstance(f, dict)]
    for g in (parse(r['gold']) or {}).get('findings', []):
        gq = Q(g)
        hits = [f for f in F if gq and sum(any(qh(x, y) for y in Q(f)) for x in gq) / len(gq) >= 0.5]
        if hits and not any(f.get('rule_id') == g.get('rule_id') for f in hits):
            rows.append((i, r['project'], g, hits[0]))
rows.sort(key=lambda x: (x[2].get('rule_id'), x[3].get('rule_id'), x[0]))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '规则贴错'
ws.merge_cells('A1:K1')
ws['A1'] = (f'开发集“位置找对、规则贴错”清单（{len(rows)} 条，占标准意见 106 条的 {len(rows)/106:.1%}）：模型意见的引文已覆盖专家引文，但规则编号与专家不同。'
            '请对每条判断：模型贴错=专家规则明显更合适；两个都行=两个规则都说得通（评分口径过严）；专家贴错=模型的规则更合适。理由写备注列。')
ws['A1'].alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[1].height = 48
H = ['序号', '样本', '项目', '专家规则', '模型规则', '专家引文', '专家意见', '模型意见', '专家规则判断要点', '判定', '备注']
for c, h in enumerate(H, 1):
    x = ws.cell(2, c, h); x.font = Font(name='Arial', bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78'); x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin = Side(style='thin', color='BBBBBB'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
for n, (i, proj, g, f) in enumerate(rows, 1):
    vals = [n, i, proj, name(g.get('rule_id')), name(f.get('rule_id')), '\n'.join(q for q in Q(g))[:400], txt(g)[:500], txt(f)[:400], pts(g.get('rule_id'))[:300], None, None]
    for c, v in enumerate(vals, 1):
        x = ws.cell(2 + n, c, v); x.font = Font(name='Arial', size=10); x.alignment = Alignment(wrap_text=True, vertical='top'); x.border = bd
for k, w in zip('ABCDEFGHIJK', [6, 16, 20, 20, 20, 40, 50, 40, 36, 12, 26]): ws.column_dimensions[k].width = w
dv = DataValidation(type='list', formula1='"模型贴错,两个都行,专家贴错"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'J3:J{2 + len(rows)}')
ws.freeze_panes = 'D3'
# 汇总页：规则对计数
ws2 = wb.create_sheet('混淆对计数')
cnt = collections.Counter((g.get('rule_id'), f.get('rule_id')) for _, _, g, f in rows)
for c, h in enumerate(['专家规则', '模型规则', '条数'], 1):
    x = ws2.cell(1, c, h); x.font = Font(name='Arial', bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78')
for k, ((a, b), v) in enumerate(cnt.most_common(), 2):
    for c, val in enumerate([name(a), name(b), v], 1): ws2.cell(k, c, val).font = Font(name='Arial', size=10)
for k, w in zip('ABC', [26, 26, 8]): ws2.column_dimensions[k].width = w
wb.save(OUT); print(OUT, len(rows))
