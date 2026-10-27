"""意见复核样本抽查表。Usage: python3 -I make_judge_xlsx.py <out_judge/samples.jsonl> <out.xlsx>"""
import sys, json, re, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
rows = [json.loads(l) for l in open(sys.argv[1], encoding='utf-8')]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '意见复核样本'
ws.append(['序号', '样本', '项目', '规则', '候选意见', '标准答案', '原因类型', '依据的报告原文', '理由', '抽查意见（同意/不同意）', '备注'])
for i, r in enumerate(rows, 1):
    o = json.loads(r['messages'][1]['content'].split('</think>')[1]); m = r['meta']
    ws.append([i, r['id'], m['project_id'], m['rule_id'], m['candidate'], m['verdict'], m['reason_type'], '\n'.join(e['quote'] for e in o['evidence']), o['reason'], None, None])
for c in ws[1]: c.font = Font(bold=True, name='Arial'); c.fill = PatternFill('solid', fgColor='DDEBF7')
for col, w in zip('ABCDEFGHIJK', (6, 20, 8, 18, 50, 10, 12, 50, 50, 14, 20)): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top'); c.font = Font(name='Arial', size=10)
dv = DataValidation(type='list', formula1='"同意,不同意"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'J2:J{len(rows) + 1}')
ws.freeze_panes = 'A2'; wb.save(sys.argv[2]); print(len(rows))
