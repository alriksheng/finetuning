"""训练规则标签复核表：候选（flag_label_noise.py 输出）+ 初判（data/label_review/label_review_decisions.py）→ xlsx。
Usage: python3 -I make_label_review_xlsx.py <data_dir> <候选.json> <初判.py> <out.xlsx>"""
import sys, os, json, runpy, collections, yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
D, CAND, DEC, OUT = sys.argv[1:5]
cand = json.load(open(CAND, encoding='utf-8')); dec = runpy.run_path(DEC)['D']
R = {r['rule_id']: r['tag'] for r in yaml.safe_load(open(os.path.join(D, 'rules/rules.yaml'), encoding='utf-8'))}
focus = {}
for l in open(os.path.join(D, 'out/samples.jsonl'), encoding='utf-8'):
    r = json.loads(l); focus[r['id']] = set(r['meta'].get('focus_rule_ids') or [])
nm = lambda r: f"{r}｜{R.get(r, '')}" if r else ''
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '标签复核'
n_chg = sum(1 for v in dec.values() if v[0] == '建议改标')
ws.merge_cells('A1:M1')
ws['A1'] = (f'训练意见规则标签复核（{len(cand)} 条）：用意见文字训练轻量分类器、按项目分组交叉验证，找出“标注规则与文字预测规则高置信度不一致”的训练意见。'
            '初判由 Claude Code 依据 8 号文、15 号文逐条给出：标注正确=保持；建议改标=按意见内容应改为建议规则；两个皆可=意见含两项或两个规则相邻，保持原标注。'
            '“建议规则在审核重点内”为否的，说明该样本的审核重点没有这个规则，不能改。复核结论列留给你填写（同意/不同意），不填则默认按初判。')
ws['A1'].alignment = Alignment(wrap_text=True, vertical='center'); ws.row_dimensions[1].height = 62
H = ['序号', '样本', '项目', '小节', '现标注规则', '文字预测规则（概率）', '意见', '引文（首条）', '初判', '建议规则', '建议规则在审核重点内', '理由', '复核结论']
for c, h in enumerate(H, 1):
    x = ws.cell(2, c, h); x.font = Font(name='Arial', bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78'); x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin = Side(style='thin', color='BBBBBB'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
fills = {'建议改标': 'FFF2CC', '标注正确': 'E2EFDA', '两个皆可': 'EDEDED'}
rows = []
for i, x in enumerate(cand, 1):
    v, new, why = dec[i]
    infoc = ''
    if v == '建议改标':
        infoc = '是' if new in focus.get(x['sid'], set()) else '否'
        if infoc == '否': why += '（该样本审核重点内没有此规则，不宜改，按“两个皆可”保留）'; v = '两个皆可'; new = None
    rows.append((i, x, v, new, infoc, why))
order = {'建议改标': 0, '标注正确': 1, '两个皆可': 2}
rows.sort(key=lambda t: (order[t[2]], t[1]['rule'], t[3] or '', t[0]))
for n, (i, x, v, new, infoc, why) in enumerate(rows, 3):
    vals = [i, x['sid'], x['project'], (x['scope'] or '')[:60], nm(x['rule']), f"{nm(x['pred'])}（{x['prob']:.2f}）", x['text'][:500], (x['quotes'][0] if x['quotes'] else '')[:160], v, nm(new), infoc, why, None]
    for c, val in enumerate(vals, 1):
        y = ws.cell(n, c, val); y.font = Font(name='Arial', size=10); y.alignment = Alignment(wrap_text=True, vertical='top'); y.border = bd
    ws.cell(n, 9).fill = PatternFill('solid', fgColor=fills[v])
for k, w in zip('ABCDEFGHIJKLM', [6, 16, 18, 26, 22, 24, 50, 36, 10, 22, 10, 50, 12]): ws.column_dimensions[k].width = w
dv = DataValidation(type='list', formula1='"同意,不同意"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'M3:M{2 + len(rows)}')
ws.freeze_panes = 'C3'
s = wb.create_sheet('汇总', 0)
last = 2 + len(rows)
hdr = ['初判', '条数', '占比', '含义']
meaning = {'建议改标': '按意见内容应改为建议规则（建议规则在审核重点内）', '标注正确': '现标注合适，文字分类器预测不准', '两个皆可': '意见含两项、规则相邻，或审核重点内无更合适规则；保持原标注'}
for c, h in enumerate(hdr, 1):
    x = s.cell(1, c, h); x.font = Font(name='Arial', bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78')
for r_, k in enumerate(['建议改标', '标注正确', '两个皆可'], 2):
    s.cell(r_, 1, k); s.cell(r_, 2, f"=COUNTIF('标签复核'!I3:I{last},A{r_})"); s.cell(r_, 3, f'=B{r_}/B5'); s.cell(r_, 4, meaning[k])
s.cell(5, 1, '合计'); s.cell(5, 2, '=SUM(B2:B4)'); s.cell(5, 3, '=B5/B5')
for r_ in range(2, 6): s.cell(r_, 3).number_format = '0.0%'
s.cell(7, 1, '建议改标的规则对（现标注→建议）').font = Font(name='Arial', bold=True)
cnt = collections.Counter((x['rule'], new) for _, x, v, new, _, _ in rows if v == '建议改标')
for r_, ((a, b), vv) in enumerate(cnt.most_common(), 8):
    s.cell(r_, 1, nm(a)); s.cell(r_, 2, nm(b)); s.cell(r_, 3, vv)
for row in s.iter_rows():
    for c in row: c.font = Font(name='Arial', size=10, bold=c.font.bold, color=c.font.color)
for k, w in zip('ABCD', [28, 28, 8, 60]): s.column_dimensions[k].width = w
wb.save(OUT)
json.dump([dict(sid=x['sid'], k=x['k'], old=x['rule'], new=new, verdict=v) for _, x, v, new, _, _ in rows if v == '建议改标'], open('data/label_review/label_review_changes.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(OUT, collections.Counter(v for _, _, v, _, _, _ in rows))
